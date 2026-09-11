from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import SignalRecord, TradeRecord
from .stats import build_stats, max_tp_hit
from .timeutil import display_dt

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")

TRADE_HEADERS = [
    "קבוצה",
    "מזהה איתות",
    "תאריך ושעה קבלת איתות",
    "סימול",
    "כיוון",
    "טווח כניסה",
    "מספר TP",
    "מחיר TP",
    "טיקט",
    "לוט",
    "תאריך ושעה כניסה",
    "מחיר כניסה",
    "סטופ",
    "סטופ הועבר לכניסה",
    "תאריך ושעה יציאה / מימוש",
    "מחיר יציאה",
    "סיבת יציאה",
    "רווח/הפסד",
    "פיפס",
    "סטטוס",
]

SIGNAL_HEADERS = [
    "מזהה איתות",
    "קבוצה",
    "תאריך ושעה קבלה",
    "סימול",
    "כיוון",
    "טווח",
    "SL",
    "סטטוס",
    "TP מקסימלי שנלקח",
    "תאריך ושעה סיום",
    "סיבת דילוג",
]

REASON_HE = {
    "TP": "טייק פרופיט",
    "SL": "סטופ לוס",
    "BREAKEVEN": "איזון (כניסה)",
    "SAFETY": "מנגנון בטיחות",
    "MANUAL": "ידני",
    "UNKNOWN": "לא ידוע",
}

STATUS_HE = {
    "open": "פתוח",
    "closed": "סגור",
    "failed": "נכשל",
    "active": "פעיל",
    "completed": "הושלם",
    "skipped": "דולג",
}


class ExcelReporter:
    def __init__(self, output_dir: Path, tz: ZoneInfo):
        self.output_dir = output_dir
        self.tz = tz
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def month_path(self, group_name: str, year: int, month: int, final: bool = False) -> Path:
        folder = self.output_dir / group_name
        folder.mkdir(parents=True, exist_ok=True)
        suffix = "_final" if final else ""
        return folder / f"{year:04d}-{month:02d}{suffix}.xlsx"

    def write_month(
        self,
        group_name: str,
        year: int,
        month: int,
        signals: list[SignalRecord],
        trades: list[TradeRecord],
        final: bool = False,
    ) -> Path:
        path = self.month_path(group_name, year, month, final=final)
        wb = Workbook()
        ws_trades = wb.active
        ws_trades.title = "עסקאות"
        self._write_trades(ws_trades, trades)

        ws_signals = wb.create_sheet("איתותים")
        self._write_signals(ws_signals, signals, trades)

        ws_stats = wb.create_sheet("סטטיסטיקה")
        self._write_stats(ws_stats, group_name, year, month, signals, trades)

        ws_daily = wb.create_sheet("סיכום יומי")
        self._write_daily(ws_daily, trades)

        for ws in wb.worksheets:
            ws.sheet_view.rightToLeft = True
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

        wb.save(path)
        return path

    def _write_trades(self, ws, trades: list[TradeRecord]) -> None:
        _header(ws, TRADE_HEADERS)
        for i, t in enumerate(trades, start=2):
            reason = t.close_reason or ""
            if t.sl_moved_to_be and (reason or "").upper() == "SL":
                reason = "BREAKEVEN"
            ws.cell(i, 1, t.group_name)
            ws.cell(i, 2, t.signal_id)
            ws.cell(i, 3, display_dt(t.received_at, self.tz))
            ws.cell(i, 4, t.symbol)
            ws.cell(i, 5, t.side)
            ws.cell(i, 6, f"{t.zone_low:g}-{t.zone_high:g}")
            ws.cell(i, 7, t.tp_index)
            ws.cell(i, 8, t.tp_price)
            ws.cell(i, 9, t.ticket)
            ws.cell(i, 10, t.lot)
            ws.cell(i, 11, display_dt(t.entry_time, self.tz))
            ws.cell(i, 12, t.entry_price)
            ws.cell(i, 13, t.sl)
            ws.cell(i, 14, "כן" if t.sl_moved_to_be else "לא")
            ws.cell(i, 15, display_dt(t.exit_time, self.tz))
            ws.cell(i, 16, t.exit_price)
            ws.cell(i, 17, REASON_HE.get(reason.upper(), reason) if reason else "")
            ws.cell(i, 18, t.profit)
            ws.cell(i, 19, t.pips)
            ws.cell(i, 20, STATUS_HE.get(t.status, t.status))
        _autosize(ws, TRADE_HEADERS)

    def _write_signals(self, ws, signals: list[SignalRecord], trades: list[TradeRecord]) -> None:
        _header(ws, SIGNAL_HEADERS)
        by_sig: dict[str, list[TradeRecord]] = defaultdict(list)
        for t in trades:
            by_sig[t.signal_id].append(t)
        for i, s in enumerate(signals, start=2):
            level = s.max_tp_hit
            if level is None:
                level = max_tp_hit(by_sig.get(s.id, []))
            ws.cell(i, 1, s.id)
            ws.cell(i, 2, s.group_name)
            ws.cell(i, 3, display_dt(s.received_at, self.tz))
            ws.cell(i, 4, s.symbol)
            ws.cell(i, 5, s.side)
            ws.cell(i, 6, f"{s.zone_low:g}-{s.zone_high:g}")
            ws.cell(i, 7, s.sl)
            ws.cell(i, 8, STATUS_HE.get(s.status, s.status))
            ws.cell(i, 9, level)
            ws.cell(i, 10, display_dt(s.completed_at, self.tz))
            ws.cell(i, 11, s.skipped_reason)
        _autosize(ws, SIGNAL_HEADERS)

    def _write_stats(
        self,
        ws,
        group_name: str,
        year: int,
        month: int,
        signals: list[SignalRecord],
        trades: list[TradeRecord],
    ) -> None:
        stats = build_stats(signals, trades)
        rows = [
            ("קבוצה", group_name),
            ("חודש", f"{year:04d}-{month:02d}"),
            ("סה״כ איתותים", stats["signals_total"]),
            ("הושלמו", stats["signals_completed"]),
            ("פעילים", stats["signals_active"]),
            ("דולגו (עסוק / כפול)", stats["signals_skipped"]),
            ("נכשלו", stats["signals_failed"]),
            ("רווח/הפסד כולל", stats["profit_total"]),
            ("ממוצע לאיתות שהושלם", stats["profit_avg"]),
            ("", ""),
            (
                "אחוזי פגיעה בטייקים (מתוך איתותים שהושלמו)",
                f"n={stats['denominator']}",
            ),
        ]
        _header(ws, ["מדד", "ערך"])
        r = 2
        for label, value in rows:
            ws.cell(r, 1, label)
            ws.cell(r, 2, value)
            r += 1
        r += 1
        ws.cell(r, 1, "TP")
        ws.cell(r, 2, "כמה איתותים הגיעו לפחות עד כאן")
        ws.cell(r, 3, "אחוז")
        for cell in (ws.cell(r, 1), ws.cell(r, 2), ws.cell(r, 3)):
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        r += 1
        for n in range(1, stats["max_tp_levels"] + 1):
            ws.cell(r, 1, f"TP{n}")
            ws.cell(r, 2, stats["hit_counts"].get(n, 0))
            ws.cell(r, 3, f"{stats['hit_pct'].get(n, 0)}%")
            r += 1
        r += 2
        ws.cell(r, 1, "התפלגות TP מקסימלי שנלקח")
        r += 1
        ws.cell(r, 1, "TP מקסימלי")
        ws.cell(r, 2, "מספר איתותים")
        ws.cell(r, 3, "אחוז")
        for cell in (ws.cell(r, 1), ws.cell(r, 2), ws.cell(r, 3)):
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        r += 1
        for level, count in sorted(stats["distribution"].items()):
            label = "יציאה לפני TP1 (סטופ/אחר)" if level == 0 else f"TP{level}"
            ws.cell(r, 1, label)
            ws.cell(r, 2, count)
            ws.cell(r, 3, f"{stats['distribution_pct'].get(level, 0)}%")
            r += 1
        _autosize(ws, ["מדד", "ערך", "אחוז"])

    def _write_daily(self, ws, trades: list[TradeRecord]) -> None:
        headers = ["תאריך", "עסקאות", "סגורות", "רווח/הפסד", "פיפס"]
        _header(ws, headers)
        buckets: dict[str, list[TradeRecord]] = defaultdict(list)
        for t in trades:
            day = display_dt(t.received_at, self.tz)[:10]
            if day:
                buckets[day].append(t)
        for i, day in enumerate(sorted(buckets), start=2):
            group = buckets[day]
            closed = [t for t in group if t.status == "closed"]
            ws.cell(i, 1, day)
            ws.cell(i, 2, len(group))
            ws.cell(i, 3, len(closed))
            ws.cell(i, 4, round(sum(t.profit or 0.0 for t in closed), 2))
            ws.cell(i, 5, round(sum(t.pips or 0.0 for t in closed), 1))
        _autosize(ws, headers)


def month_bounds(year: int, month: int, tz: ZoneInfo) -> tuple[datetime, datetime]:
    start = datetime(year, month, 1, tzinfo=tz)
    if month == 12:
        end = datetime(year + 1, 1, 1, tzinfo=tz)
    else:
        end = datetime(year, month + 1, 1, tzinfo=tz)
    return start, end


def _header(ws, headers: list[str]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP


def _autosize(ws, headers: list[str]) -> None:
    for i, name in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(36, len(name) + 4))
