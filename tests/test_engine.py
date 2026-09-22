from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from src.brokers.dry_run import DryRunBroker
from src.config import AppConfig
from src.db import Database
from src.excel_report import ExcelReporter
from src.models import CloseReason, GroupConfig, Side
from src.parser import parse_signal
from src.trade_manager import IncomingSignal, TradingEngine

SELL = """#XAUUSD SELL 4315-4318

TP 4312
TP 4309
TP 4305
TP 4300
TP 4295
TP 4285

SL 4327
"""

BUY = """#XAUUSD BUY 4349-4346

TP 4352
TP 4355
TP 4359
TP 4364
TP 4369
TP 4379

SL 4337
"""


def _cfg(tmp: Path) -> AppConfig:
    tz = ZoneInfo("Asia/Jerusalem")
    return AppConfig(
        timezone="Asia/Jerusalem",
        dry_run=True,
        session_name="test",
        groups=[
            GroupConfig("group_a", -1001, 0.01),
            GroupConfig("group_b", -1002, 0.01),
        ],
        admin_chat_id=0,
        telegram_api_id=0,
        telegram_api_hash="",
        telegram_phone="",
        symbol_override="",
        max_concurrent_signals=1,
        skip_first_tps=2,
        breakeven_after_tp=3,
        safety_delay_seconds=5,
        safety_pips=50,
        guard_interval_seconds=5,
        pip_size=0.1,
        magic_number=260908,
        deviation=30,
        order_retries=0,
        broker_type="dry_run",
        common_files_dir="",
        mt5_login=None,
        mt5_password="",
        mt5_server="",
        mt5_terminal_path="",
        output_dir=tmp / "reports",
        daily_digest_hour=23,
        daily_digest_minute=55,
        data_dir=tmp / "data",
        logs_dir=tmp / "logs",
        lot_size_default=0.01,
        tz=tz,
    )


def _engine(tmp: Path) -> tuple[TradingEngine, Database, DryRunBroker]:
    cfg = _cfg(tmp)
    cfg.data_dir.mkdir(parents=True)
    cfg.output_dir.mkdir(parents=True)
    db = Database(cfg.data_dir / "bot.db")
    broker = DryRunBroker()
    broker.connect()
    reporter = ExcelReporter(cfg.output_dir, cfg.tz)
    engine = TradingEngine(cfg, db, broker, reporter)
    return engine, db, broker


def _incoming(text: str, group: GroupConfig, msg_id: int) -> IncomingSignal:
    parsed = parse_signal(text)
    assert parsed is not None
    return IncomingSignal(
        parsed=parsed,
        group=group,
        chat_id=group.chat_id,
        telegram_msg_id=msg_id,
        received_at=datetime.now(timezone.utc),
        raw_text=text,
    )


def test_execute_then_persist(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    group = engine.cfg.groups[0]
    engine._handle(_incoming(SELL, group, 11))
    signals = db.active_signals()
    assert len(signals) == 1
    trades = db.trades_for_signal(signals[0].id)
    # שתי שורות ה-TP הראשונות מדולגות, המספור נשאר לפי השורה במקור.
    assert len(trades) == 4
    assert all(t.status == "open" and t.ticket for t in trades)
    assert [t.tp_index for t in trades] == [3, 4, 5, 6]
    assert [t.tp_price for t in trades] == [4305.0, 4300.0, 4295.0, 4285.0]
    assert broker.positions(engine.cfg.magic_number)
    path = engine._write_excel(group.name, datetime.now(timezone.utc), final=False)
    assert path and Path(path).exists()


def test_one_signal_at_a_time(tmp_path: Path):
    engine, db, _broker = _engine(tmp_path)
    a, b = engine.cfg.groups
    engine._handle(_incoming(SELL, a, 1))
    engine._handle(_incoming(BUY, b, 2))
    skipped = [s for s in db.signals_in_range(b.name, datetime(2020, 1, 1, tzinfo=timezone.utc), datetime(2030, 1, 1, tzinfo=timezone.utc)) if s.status == "skipped"]
    assert len(skipped) == 1
    assert skipped[0].skipped_reason == "busy_one_signal"
    assert len(db.active_signals()) == 1


def test_tp3_moves_remaining_stops_to_entry(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    group = engine.cfg.groups[0]
    engine._handle(_incoming(SELL, group, 3))
    sid = db.active_signals()[0].id
    trades = db.trades_for_signal(sid)
    for trade in trades:
        if trade.tp_index <= 3 and trade.ticket:
            broker.simulate_close(trade.ticket, CloseReason.TP, exit_price=trade.tp_price, profit=8)
    engine._monitor()
    remaining = [t for t in db.trades_for_signal(sid) if t.status == "open"]
    assert len(remaining) == 3
    assert [t.tp_index for t in remaining] == [4, 5, 6]
    assert all(t.sl_moved_to_be for t in remaining)
    assert all(t.sl == t.entry_price for t in remaining)


def test_safety_fills_missing_sl_tp(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    group = engine.cfg.groups[0]
    engine._handle(_incoming(SELL, group, 4))
    trade = db.open_trades()[0]
    assert trade.ticket is not None
    broker.modify_sl_tp(trade.ticket, sl=0.0, tp=0.0)
    engine._apply_safety(trade.signal_id)
    pos = {p.ticket: p for p in broker.positions(engine.cfg.magic_number)}[trade.ticket]
    assert pos.sl != 0
    assert pos.tp != 0
    delta = engine.cfg.safety_pips * engine.cfg.pip_size
    assert abs(pos.sl - (trade.entry_price + delta)) < 1e-6
    assert abs(pos.tp - (trade.entry_price - delta)) < 1e-6


def test_excel_files_are_isolated_per_group(tmp_path: Path):
    from openpyxl import load_workbook

    engine, db, broker = _engine(tmp_path)
    group_a, group_b = engine.cfg.groups
    engine._handle(_incoming(SELL, group_a, 10))
    for trade in db.open_trades():
        if trade.ticket:
            broker.simulate_close(trade.ticket, CloseReason.TP, exit_price=trade.tp_price, profit=1)
    engine._monitor()
    engine._handle(_incoming(BUY, group_b, 20))

    now = datetime.now(timezone.utc)
    path_a = Path(engine._write_excel(group_a.name, now, final=False))
    path_b = Path(engine._write_excel(group_b.name, now, final=False))
    assert path_a.parent.name == "group_a"
    assert path_b.parent.name == "group_b"
    assert path_a != path_b

    book_a = load_workbook(path_a)["עסקאות"]
    book_b = load_workbook(path_b)["עסקאות"]
    names_a = {book_a.cell(row, 1).value for row in range(2, book_a.max_row + 1)}
    names_b = {book_b.cell(row, 1).value for row in range(2, book_b.max_row + 1)}
    assert names_a == {"group_a"}
    assert names_b == {"group_b"}
    assert book_a.max_row == 5
    assert book_b.max_row == 5


def test_skips_signal_when_all_tp_lines_are_skipped(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    group = engine.cfg.groups[0]
    short = "#XAUUSD SELL 4315-4318\nTP 4312\nTP 4309\nSL 4327\n"
    engine._handle(_incoming(short, group, 77))
    assert db.active_signals() == []
    assert broker.positions(engine.cfg.magic_number) == []
    window = (datetime(2020, 1, 1, tzinfo=timezone.utc), datetime(2030, 1, 1, tzinfo=timezone.utc))
    signals = db.signals_in_range(group.name, *window)
    assert [s.skipped_reason for s in signals] == ["no_tp_after_skip"]


def test_guard_forces_sl_tp_on_any_position(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    group = engine.cfg.groups[0]
    engine._handle(_incoming(SELL, group, 5))
    for trade in db.open_trades():
        broker.modify_sl_tp(trade.ticket, sl=0.0, tp=0.0)

    engine._guard_positions()

    for pos in broker.positions(engine.cfg.magic_number):
        assert pos.sl != 0 and pos.tp != 0


def test_guard_protects_position_missing_from_db(tmp_path: Path):
    engine, _db, broker = _engine(tmp_path)
    # פוזיציה שנפתחה אך לא הספיקה להיכתב ל-DB לפני קריסה.
    result = broker.market_order(
        symbol="XAUUSD",
        side=Side.SELL,
        volume=0.01,
        sl=0.0,
        tp=0.0,
        comment="GS|group_a|T4",
        deviation=30,
        magic=engine.cfg.magic_number,
    )
    engine._guard_positions()
    pos = {p.ticket: p for p in broker.positions(engine.cfg.magic_number)}[result.ticket]
    delta = engine.cfg.safety_pips * engine.cfg.pip_size
    assert abs(pos.sl - (pos.price_open + delta)) < 1e-6
    assert abs(pos.tp - (pos.price_open - delta)) < 1e-6


def test_recovery_closes_trades_that_ended_while_down(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    group = engine.cfg.groups[0]
    engine._handle(_incoming(SELL, group, 6))
    gone = db.open_trades()[0]
    broker.simulate_close(gone.ticket, CloseReason.TP, exit_price=gone.tp_price, profit=12)

    # מדמים הפעלה מחדש: אותו DB, מנוע חדש.
    engine._be_done.clear()
    engine._recover()

    restored = db.trade_by_ticket(gone.ticket)
    assert restored.status == "closed"
    assert restored.close_reason == CloseReason.TP.value
    assert restored.profit == 12


PLAN_SELL = """XAUUSD SELL PLAN @4365_4368

tp @4349
tp @4335

SL@4381
"""


def test_plan_template_group_enters_every_tp(tmp_path: Path):
    engine, db, _broker = _engine(tmp_path)
    plan_group = GroupConfig(
        "plan_group", -1003, 0.01, magic=260910, skip_first_tps=0, breakeven_after_tp=1
    )
    engine.cfg.groups.append(plan_group)

    engine._handle(_incoming(PLAN_SELL, plan_group, 90))

    sid = db.active_signals()[0].id
    trades = db.trades_for_signal(sid)
    assert [t.tp_index for t in trades] == [1, 2]
    assert [t.tp_price for t in trades] == [4349.0, 4335.0]
    assert all(t.sl == 4381.0 for t in trades)


def test_plan_group_moves_stops_after_first_tp(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    plan_group = GroupConfig(
        "plan_group", -1003, 0.01, magic=260910, skip_first_tps=0, breakeven_after_tp=1
    )
    engine.cfg.groups.append(plan_group)
    engine._handle(_incoming(PLAN_SELL, plan_group, 91))
    sid = db.active_signals()[0].id

    first = next(t for t in db.trades_for_signal(sid) if t.tp_index == 1)
    broker.simulate_close(first.ticket, CloseReason.TP, exit_price=first.tp_price, profit=5)
    engine._monitor()

    remaining = [t for t in db.trades_for_signal(sid) if t.status == "open"]
    assert len(remaining) == 1
    assert remaining[0].tp_index == 2
    assert remaining[0].sl_moved_to_be
    assert remaining[0].sl == remaining[0].entry_price


def test_two_groups_keep_separate_strategies(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    skip_group = engine.cfg.groups[0]
    plan_group = GroupConfig(
        "plan_group", -1003, 0.01, magic=260910, skip_first_tps=0, breakeven_after_tp=1
    )
    engine.cfg.groups.append(plan_group)

    engine._handle(_incoming(SELL, skip_group, 92))
    sid_a = db.active_signals()[0].id
    for trade in db.trades_for_signal(sid_a):
        broker.simulate_close(trade.ticket, CloseReason.TP, exit_price=trade.tp_price, profit=1)
    engine._monitor()

    engine._handle(_incoming(PLAN_SELL, plan_group, 93))
    sid_b = db.active_signals()[0].id

    assert [t.tp_index for t in db.trades_for_signal(sid_a)] == [3, 4, 5, 6]
    assert [t.tp_index for t in db.trades_for_signal(sid_b)] == [1, 2]


def _production_groups() -> list[GroupConfig]:
    """בדיוק מה שנכנס ל-config.yaml בשרת."""
    return [
        GroupConfig("group_a", -1002001216034, 0.01, magic=260908,
                    skip_first_tps=2, breakeven_after_tp=3),
        GroupConfig("TechnicalPips6273", -1001569906975, 0.01,
                    username="TechnicalPips6273", magic=260909,
                    skip_first_tps=2, breakeven_after_tp=3),
        GroupConfig("profit_kings", -1002984909577, 0.01,
                    username="profitkingscalper_007", magic=260910,
                    skip_first_tps=0, breakeven_after_tp=1),
    ]


def test_production_config_end_to_end(tmp_path: Path):
    """שלוש הקבוצות, שתי התבניות, שתי האסטרטגיות, ודוח נפרד לכל קבוצה."""
    from openpyxl import load_workbook

    engine, db, broker = _engine(tmp_path)
    engine.cfg.groups = _production_groups()
    vip, pips, kings = engine.cfg.groups

    # אסטרטגיה א' על התבנית הישנה: 6 שורות TP -> 4 פוזיציות, החל מ-TP3.
    engine._handle(_incoming(SELL, vip, 101))
    sid_vip = db.active_signals()[0].id
    vip_trades = db.trades_for_signal(sid_vip)
    assert [t.tp_index for t in vip_trades] == [3, 4, 5, 6]
    assert all(t.sl == 4327.0 for t in vip_trades)
    assert all(t.side == "SELL" for t in vip_trades)

    # נגיעה ב-TP3 מקדמת את כל הסטופים לנקודת הכניסה.
    tp3 = next(t for t in vip_trades if t.tp_index == 3)
    broker.simulate_close(tp3.ticket, CloseReason.TP, exit_price=tp3.tp_price, profit=4)
    engine._monitor()
    still_open = [t for t in db.trades_for_signal(sid_vip) if t.status == "open"]
    assert [t.tp_index for t in still_open] == [4, 5, 6]
    assert all(t.sl_moved_to_be and t.sl == t.entry_price for t in still_open)

    for trade in still_open:
        broker.simulate_close(trade.ticket, CloseReason.TP, exit_price=trade.tp_price, profit=6)
    engine._monitor()

    # אסטרטגיה ב' על תבנית ה-PLAN: 2 שורות TP -> 2 פוזיציות, החל מ-TP1.
    engine._handle(_incoming(PLAN_SELL, kings, 102))
    sid_kings = db.active_signals()[0].id
    king_trades = db.trades_for_signal(sid_kings)
    assert [t.tp_index for t in king_trades] == [1, 2]
    assert all(t.sl == 4381.0 for t in king_trades)

    first = next(t for t in king_trades if t.tp_index == 1)
    broker.simulate_close(first.ticket, CloseReason.TP, exit_price=first.tp_price, profit=3)
    engine._monitor()
    king_open = [t for t in db.trades_for_signal(sid_kings) if t.status == "open"]
    assert [t.tp_index for t in king_open] == [2]
    assert king_open[0].sl_moved_to_be

    # קבוצה שלישית שלא נסחרה בכלל לא מייצרת רעש בדוחות.
    now = datetime.now(timezone.utc)
    paths = {g.name: Path(engine._write_excel(g.name, now, final=False)) for g in (vip, pips, kings)}
    assert len({p.resolve() for p in paths.values()}) == 3
    assert paths["group_a"].parent.name == "group_a"
    assert paths["profit_kings"].parent.name == "profit_kings"

    sheet_vip = load_workbook(paths["group_a"])["עסקאות"]
    sheet_kings = load_workbook(paths["profit_kings"])["עסקאות"]
    assert {sheet_vip.cell(r, 1).value for r in range(2, sheet_vip.max_row + 1)} == {"group_a"}
    assert {sheet_kings.cell(r, 1).value for r in range(2, sheet_kings.max_row + 1)} == {"profit_kings"}
    assert sheet_vip.max_row == 5      # 4 עסקאות + כותרת
    assert sheet_kings.max_row == 3    # 2 עסקאות + כותרת

    # קבוצה בלי פעילות מקבלת קובץ עם כותרת בלבד.
    sheet_pips = load_workbook(paths["TechnicalPips6273"])["עסקאות"]
    assert sheet_pips.max_row == 1


def test_data_survives_restart(tmp_path: Path):
    """אחרי קריסה, המידע נשאר ב-DB והדוח נבנה מחדש ממנו."""
    from openpyxl import load_workbook

    engine, db, broker = _engine(tmp_path)
    engine.cfg.groups = _production_groups()
    vip = engine.cfg.groups[0]
    engine._handle(_incoming(SELL, vip, 201))

    # העסקאות נסגרות בזמן שהבוט למטה.
    for trade in db.open_trades():
        broker.simulate_close(trade.ticket, CloseReason.TP, exit_price=trade.tp_price, profit=7)

    engine._recover()

    closed = [t for t in db.trades_for_signal(db.signals_in_range(
        vip.name,
        datetime(2020, 1, 1, tzinfo=timezone.utc),
        datetime(2030, 1, 1, tzinfo=timezone.utc),
    )[0].id)]
    assert len(closed) == 4
    assert all(t.status == "closed" and t.profit == 7 for t in closed)

    path = Path(engine._write_excel(vip.name, datetime.now(timezone.utc), final=False))
    sheet = load_workbook(path)["עסקאות"]
    assert sheet.max_row == 5


def test_recovery_adopts_orphan_position(tmp_path: Path):
    engine, db, broker = _engine(tmp_path)
    result = broker.market_order(
        symbol="XAUUSD",
        side=Side.BUY,
        volume=0.01,
        sl=4300.0,
        tp=4360.0,
        comment="GS|group_b|T5",
        deviation=30,
        magic=engine.cfg.magic_number,
    )
    engine._recover()

    adopted = db.trade_by_ticket(result.ticket)
    assert adopted is not None
    assert adopted.status == "open"
    assert adopted.group_name == "group_b"
    assert adopted.tp_index == 5
    assert adopted.entry_price == result.fill_price
