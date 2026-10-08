from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from src.excel_report import ExcelReporter
from src.models import SignalRecord, TradeRecord
from src.timeutil import utcnow

SELL_TEXT = """#XAUUSD SELL 4315-4318

TP 4312
TP 4309
TP 4305
TP 4300
TP 4295
TP 4285

SL 4327
"""


def _signal(**kw) -> SignalRecord:
    now = utcnow()
    data = dict(
        id="sig-1",
        group_name="ChannelB",
        chat_id=-1009876543210,
        telegram_msg_id=11,
        received_at=now,
        symbol="XAUUSD",
        side="SELL",
        zone_low=4315,
        zone_high=4318,
        sl=4327,
        raw_text=SELL_TEXT,
        status="active",
        username="ChannelB",
    )
    data.update(kw)
    return SignalRecord(**data)


def _trade(tp_index: int, **kw) -> TradeRecord:
    now = datetime(2026, 10, 2, 10, 15, 3, tzinfo=timezone.utc)
    data = dict(
        id=tp_index,
        ticket=1000 + tp_index,
        signal_id="sig-1",
        group_name="ChannelB",
        tp_index=tp_index,
        tp_price=4300.0 if tp_index == 4 else 4295.0 if tp_index == 5 else 4285.0,
        lot=0.01,
        side="SELL",
        symbol="XAUUSD",
        received_at=now,
        entry_time=now,
        entry_price=4316.2,
        sl=4327.0,
        original_sl=4327.0,
        sl_moved_to_be=False,
        exit_time=None,
        exit_price=None,
        close_reason=None,
        profit=None,
        pips=None,
        status="open",
        zone_low=4315,
        zone_high=4318,
        telegram_username="ChannelB",
    )
    data.update(kw)
    return TradeRecord(**data)


def test_excel_filename_includes_group(tmp_path: Path):
    reporter = ExcelReporter(tmp_path, ZoneInfo("Asia/Jerusalem"))
    path = reporter.write_month("ChannelB", 2026, 10, [_signal()], [_trade(4)])
    assert path.name == "ChannelB_2026-10.xlsx"
    assert path.parent.name == "ChannelB"


def test_excel_detail_is_one_to_one_with_signal_tps(tmp_path: Path):
    reporter = ExcelReporter(tmp_path, ZoneInfo("Asia/Jerusalem"))
    closed = _trade(
        4,
        status="closed",
        close_reason="TP",
        exit_time=datetime(2026, 10, 2, 11, 2, 11, tzinfo=timezone.utc),
        exit_price=4300.0,
        profit=4.2,
    )
    open_leg = _trade(5)
    path = reporter.write_month(
        "ChannelB",
        2026,
        10,
        [_signal()],
        [closed, open_leg],
    )
    book = load_workbook(path)

    trades = book["עסקאות"]
    assert trades.cell(1, 11).value == "סדר בין הנלקחות"
    assert trades.cell(1, 12).value == "האם ה-TP ניקלח"
    assert trades.cell(1, 15).value == "תאריך ושעה כניסה"
    assert trades.cell(1, 19).value == "תאריך ושעה יציאה / מימוש"
    ranks = {trades.cell(r, 9).value: trades.cell(r, 11).value for r in range(2, trades.max_row + 1)}
    assert ranks == {4: 1, 5: 2}
    hits = {trades.cell(r, 9).value: trades.cell(r, 12).value for r in range(2, trades.max_row + 1)}
    assert hits[4] == "כן"
    assert hits[5] == "לא (פתוח)"
    assert "2026-10-02" in str(trades.cell(2, 15).value)

    detail = book["פירוט איתות"]
    assert detail.cell(2, 8).value == 2
    summary = detail.cell(2, 9).value
    assert "נלקחו TP4, TP5" in summary
    assert "ניקלחו TP4" in summary
    assert "לא נלקחו TP1, TP2, TP3, TP6" in summary
    assert "TP1: לא נלקח" in str(detail.cell(2, 12).value)
    assert "ניקלח: כן" in str(detail.cell(2, 15).value)
    assert "ניקלח: לא (עדיין פתוח)" in str(detail.cell(2, 16).value)
    assert "כניסה:" in str(detail.cell(2, 15).value)
    assert "יציאה:" in str(detail.cell(2, 15).value)

    stats = book["סטטיסטיקה"]
    labels = {stats.cell(r, 1).value: stats.cell(r, 2).value for r in range(2, 20)}
    assert labels["עסקאות בגיליון (1:1 מול הדאטה)"] == 2
    assert labels["פוזיציות שנלקחו"] == 2
    assert labels["נסגרו בטייק"] == 1
